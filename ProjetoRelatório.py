import json
import logging
import os
import subprocess
import sys
import tkinter as tk
from datetime import datetime
from tkinter import filedialog, messagebox, ttk
from docx import Document
import pandas as pd

if getattr(sys, "frozen", False):
  PASTA_BASE = os.path.dirname(sys.executable)
else:
  PASTA_BASE = os.path.dirname(os.path.abspath(__file__))

LOG_FILE = os.path.join(PASTA_BASE, "gerador.log")
logging.basicConfig(
    filename=LOG_FILE,
    level=logging.INFO,
    format="%(asctime)s - [%(levelname)s] - %(message)s",
    encoding="utf-8",
)

CONFIG_FILE = os.path.join(PASTA_BASE, "config.json")


def carregar_config():
  config_padrao = {
      "pasta_destino": os.path.join(PASTA_BASE, "relatorios_gerados"),
      "modelo_template": "",
      "exportar_pdf": False,
  }
  if os.path.exists(CONFIG_FILE):
    try:
      with open(CONFIG_FILE, "r", encoding="utf-8") as f:
        config_padrao.update(json.load(f))
      logging.info("Configurações carregadas com sucesso.")
    except Exception as e:
      logging.error(f"Erro ao carregar config.json: {e}")
  return config_padrao


def salvar_config(config_dict):
  try:
    with open(CONFIG_FILE, "w", encoding="utf-8") as f:
      json.dump(config_dict, f, indent=4, ensure_ascii=False)
    logging.info("Configurações salvas no config.json.")
  except Exception as e:
    logging.error(f"Erro ao salvar config.json: {e}")


def converter_para_pdf(caminho_docx):
  try:
    caminho_docx_abs = os.path.abspath(caminho_docx).replace("\\", "/")
    caminho_pdf_abs = caminho_docx_abs.replace(".docx", ".pdf")

    ps_cmd = (
        f'$word = New-Object -ComObject Word.Application; '
        f'$word.Visible = $False; '
        f'$doc = $word.Documents.Open("{caminho_docx_abs}"); '
        f'$doc.SaveAs("{caminho_pdf_abs}", 17); '
        f'$doc.Close(); '
        f'$word.Quit()'
    )

    subprocess.run(
        ["powershell", "-Command", ps_cmd], check=True, creationflags=0x08000000
    )
    logging.info(f"PDF gerado com sucesso via PowerShell: {caminho_pdf_abs}")
    return True
  except Exception as e:
    logging.error(f"Erro ao converter para PDF (Verifique se o MS Word está instalado): {e}")
    return False


def criar_documento_word(
    titulo,
    autor,
    resumo,
    itens=None,
    identificador="",
    template_path="",
    exportar_pdf=False,
    pasta_destino="",
):
  if not pasta_destino:
    pasta_destino = os.path.join(PASTA_BASE, "relatorios_gerados")
  os.makedirs(pasta_destino, exist_ok=True)

  if template_path and os.path.exists(template_path):
    try:
      doc = Document(template_path)
      logging.info(f"Utilizando template personalizado: {template_path}")
    except Exception as e:
      logging.error(f"Falha ao carregar template ({e}). Criando documento em branco.")
      doc = Document()
  else:
    doc = Document()

  doc.add_heading(titulo, level=1)

  data_hora = datetime.now().strftime("%d/%m/%Y às %H:%M")
  p_meta = doc.add_paragraph()
  p_meta.add_run(f"Data de Geração: {data_hora}\n").italic = True
  p_meta.add_run(f"Autor / Responsável: {autor}").bold = True

  doc.add_heading("Resumo das Atividades", level=2)
  doc.add_paragraph(resumo)

  if itens:
    doc.add_heading("Itens Adicionados", level=2)
    for item in itens:
      doc.add_paragraph(item, style="List Bullet")

  nome_base = f"relatorio_{identificador}" if identificador else "relatorio"
  caminho = os.path.join(pasta_destino, f"{nome_base}.docx")

  contador = 1
  while os.path.exists(caminho):
    caminho = os.path.join(pasta_destino, f"{nome_base}_({contador}).docx")
    contador += 1

  doc.save(caminho)
  logging.info(f"Relatório Word salvo: {caminho}")

  if exportar_pdf:
    converter_para_pdf(caminho)

  return caminho


def criar_documento_consolidado(
    dados, template_path="", exportar_pdf=False, pasta_destino=""
):
  if not pasta_destino:
    pasta_destino = os.path.join(PASTA_BASE, "relatorios_gerados")
  os.makedirs(pasta_destino, exist_ok=True)

  if template_path and os.path.exists(template_path):
    try:
      doc = Document(template_path)
    except Exception:
      doc = Document()
  else:
    doc = Document()

  doc.add_heading("Relatório Consolidado de Atividades", level=1)

  data_hora = datetime.now().strftime("%d/%m/%Y às %H:%M")
  doc.add_paragraph(
      f"Gerado em: {data_hora} | Total de Registros: {len(dados)}"
  )

  for index, linha in dados.iterrows():
    nome = (
        linha.get("Nome")
        or linha.get("nome")
        or linha.get("Titulo")
        or f"Registro {index + 1}"
    )
    setor = (
        linha.get("Setor")
        or linha.get("setor")
        or linha.get("Autor")
        or "Não informado"
    )
    resumo = (
        linha.get("Resumo")
        or linha.get("resumo")
        or "Sem resumo informado."
    )

    doc.add_heading(f"#{index + 1} - {nome}", level=2)
    doc.add_paragraph(f"Setor/Autor: {setor}")
    doc.add_paragraph(resumo)
    doc.add_paragraph("-" * 50)

  caminho = os.path.join(pasta_destino, "Relatorio_Consolidado.docx")
  doc.save(caminho)
  logging.info(f"Relatório Consolidado salvo: {caminho}")

  if exportar_pdf:
    converter_para_pdf(caminho)

  return caminho


def importar_e_processar():
  caminho_arquivo = filedialog.askopenfilename(
      title="Selecione a planilha",
      filetypes=[("Planilhas Excel/CSV", "*.xlsx *.csv")],
  )
  if not caminho_arquivo:
    return

  try:
    logging.info(f"Iniciando importação de planilha: {caminho_arquivo}")
    if caminho_arquivo.endswith(".csv"):
      try:
        dados = pd.read_csv(caminho_arquivo, sep=";", encoding="latin1")
      except Exception:
        dados = pd.read_csv(caminho_arquivo)
    else:
      dados = pd.read_excel(caminho_arquivo)

    dados = dados.fillna("")
    total = len(dados)

    opcao = messagebox.askyesnocancel(
        "Modo de Geração",
        f"A planilha contém {total} registros.\n\n"
        "• Clique em SIM para gerar 1 arquivo INDIVIDUAL por linha.\n"
        "• Clique em NÃO para gerar 1 arquivo CONSOLIDADO com tudo.\n"
        "• Clique em CANCELAR para sair.",
    )

    if opcao is None:
      return

    barra_progresso["maximum"] = total
    dest = config.get("pasta_destino")
    tpl = config.get("modelo_template")
    pdf_opt = var_exportar_pdf.get()

    if opcao is True:
      for i, linha in dados.iterrows():
        nome = (
            linha.get("Nome")
            or linha.get("nome")
            or linha.get("Titulo")
            or f"Relatorio_{i+1}"
        )
        setor = (
            linha.get("Setor")
            or linha.get("setor")
            or linha.get("Autor")
            or "N/I"
        )
        resumo = (
            linha.get("Resumo")
            or linha.get("resumo")
            or f"Relatório do colaborador {nome}."
        )

        criar_documento_word(
            titulo=f"Relatório de {nome}",
            autor=setor,
            resumo=resumo,
            identificador=f"linha_{i+1}",
            template_path=tpl,
            exportar_pdf=pdf_opt,
            pasta_destino=dest,
        )

        barra_progresso["value"] = i + 1
        lbl_status.config(text=f"Gerando arquivo {i+1} de {total}...")
        janela.update()

      messagebox.showinfo(
          "Sucesso!",
          f"{total} relatórios processados com sucesso na pasta:\n'{dest}'",
      )

    else:
      lbl_status.config(text="Gerando relatório consolidado...")
      janela.update()
      caminho_final = criar_documento_consolidado(
          dados, template_path=tpl, exportar_pdf=pdf_opt, pasta_destino=dest
      )
      messagebox.showinfo(
          "Sucesso!", f"Relatório consolidado gerado em:\n{caminho_final}"
      )

  except Exception as e:
    logging.error(f"Erro no processamento em lote: {e}")
    messagebox.showerror(
        "Erro", f"Ocorreu um erro ao processar a planilha:\n{str(e)}"
    )

  finally:
    barra_progresso["value"] = 0
    lbl_status.config(text="")


def gerar_manual():
  titulo = entry_titulo.get().strip()
  autor = entry_autor.get().strip()
  resumo = text_resumo.get("1.0", tk.END).strip()
  itens = lista_itens.get(0, tk.END)

  if not titulo:
    messagebox.showwarning(
        "Atenção", "Por favor, informe pelo menos o Título do Relatório!"
    )
    return

  dest = config.get("pasta_destino")
  tpl = config.get("modelo_template")
  pdf_opt = var_exportar_pdf.get()

  caminho = criar_documento_word(
      titulo,
      autor,
      resumo,
      list(itens),
      template_path=tpl,
      exportar_pdf=pdf_opt,
      pasta_destino=dest,
  )
  messagebox.showinfo(
      "Sucesso!", f"Relatório manual gerado com sucesso!\nSalvo em: {caminho}"
  )


def selecionar_pasta_destino():
  pasta = filedialog.askdirectory(title="Selecione a Pasta de Destino")
  if pasta:
    config["pasta_destino"] = pasta
    salvar_config(config)
    lbl_pasta_destino.config(text=f"Destino: {pasta}")


def selecionar_template():
  template = filedialog.askopenfilename(
      title="Selecione o Modelo Word (.docx)",
      filetypes=[("Documentos Word", "*.docx")],
  )
  if template:
    config["modelo_template"] = template
    salvar_config(config)
    lbl_template.config(
        text=f"Modelo: {os.path.basename(template)}"
    )


def limpar_template():
  config["modelo_template"] = ""
  salvar_config(config)
  lbl_template.config(text="Modelo: Nenhum (Padrão em branco)")


def atualizar_pdf_config():
  config["exportar_pdf"] = var_exportar_pdf.get()
  salvar_config(config)


def adicionar_item():
  texto = entry_item.get().strip()
  if texto:
    lista_itens.insert(tk.END, texto)
    entry_item.delete(0, tk.END)


def remover_item():
  selecao = lista_itens.curselection()
  if selecao:
    lista_itens.delete(selecao)


config = carregar_config()

janela = tk.Tk()
janela.title("Gerador de Relatórios Word - V5.0 Pro")
janela.geometry("580x820")
janela.configure(bg="#1e1e1e")
janela.resizable(False, False)

janela.update_idletasks()
x = (janela.winfo_screenwidth() // 2) - (580 // 2)
y = (janela.winfo_screenheight() // 2) - (820 // 2)
janela.geometry(f"+{x}+{y}")

BG_DARK = "#1e1e1e"
BG_FIELD = "#2d2d2d"
FG_WHITE = "#ffffff"
FONT_TITLE = ("Segoe UI", 12, "bold")
FONT_LABEL = ("Segoe UI", 9, "bold")
FONT_MAIN = ("Segoe UI", 9)
FONT_SMALL = ("Segoe UI", 8)

style = ttk.Style()
style.theme_use("default")
style.configure(
    "Custom.Horizontal.TProgressbar",
    troughcolor="#2d2d2d",
    background="#007acc",
    thickness=10,
)

tk.Label(
    janela,
    text="GERADOR DE RELATÓRIOS V5.0",
    font=FONT_TITLE,
    bg=BG_DARK,
    fg="#007acc",
    pady=10,
).pack()

frame_config = tk.Frame(janela, bg="#252526", bd=1, relief="solid")
frame_config.pack(padx=20, fill="x", pady=5)

lbl_pasta_destino = tk.Label(
    frame_config,
    text=f"Destino: {config.get('pasta_destino')}",
    font=FONT_SMALL,
    bg="#252526",
    fg="#cccccc",
    anchor="w",
)
lbl_pasta_destino.pack(fill="x", padx=10, pady=(5, 2))

btn_pasta = tk.Button(
    frame_config,
    text="Alterar Pasta de Destino",
    command=selecionar_pasta_destino,
    bg="#3c3c3c",
    fg="white",
    bd=0,
    font=FONT_SMALL,
    cursor="hand2",
)
btn_pasta.pack(anchor="w", padx=10, pady=(0, 5))

tpl_nome = (
    os.path.basename(config.get("modelo_template"))
    if config.get("modelo_template")
    else "Nenhum (Padrão em branco)"
)
lbl_template = tk.Label(
    frame_config,
    text=f"Modelo: {tpl_nome}",
    font=FONT_SMALL,
    bg="#252526",
    fg="#cccccc",
    anchor="w",
)
lbl_template.pack(fill="x", padx=10, pady=(5, 2))

frame_tpl_btns = tk.Frame(frame_config, bg="#252526")
frame_tpl_btns.pack(anchor="w", padx=10, pady=(0, 5))

tk.Button(
    frame_tpl_btns,
    text="Selecionar Modelo (.docx)",
    command=selecionar_template,
    bg="#3c3c3c",
    fg="white",
    bd=0,
    font=FONT_SMALL,
    cursor="hand2",
).pack(side="left")
tk.Button(
    frame_tpl_btns,
    text="Usar Padrão",
    command=limpar_template,
    bg="#555555",
    fg="white",
    bd=0,
    font=FONT_SMALL,
    cursor="hand2",
).pack(side="left", padx=5)

var_exportar_pdf = tk.BooleanVar(value=config.get("exportar_pdf", False))
chk_pdf = tk.Checkbutton(
    frame_config,
    text="Exportar também em PDF (Sem necessidade de internet)",
    variable=var_exportar_pdf,
    command=atualizar_pdf_config,
    bg="#252526",
    fg="white",
    selectcolor=BG_DARK,
    activebackground="#252526",
    activeforeground="white",
    font=FONT_SMALL,
)
chk_pdf.pack(anchor="w", padx=10, pady=5)

btn_importar = tk.Button(
    janela,
    text="Importar Planilha Excel/CSV",
    command=importar_e_processar,
    bg="#2e8b57",
    fg="white",
    font=FONT_TITLE,
    bd=0,
    padx=15,
    pady=6,
    cursor="hand2",
)
btn_importar.pack(pady=10)

tk.Label(
    janela, text="──────── OU FORMULÁRIO MANUAL ────────", bg=BG_DARK, fg="#666666"
).pack(pady=5)

frame_form = tk.Frame(janela, bg=BG_DARK)
frame_form.pack(padx=25, fill="x")

tk.Label(
    frame_form,
    text="Título do Relatório:",
    font=FONT_LABEL,
    bg=BG_DARK,
    fg=FG_WHITE,
    anchor="w",
).pack(fill="x")
entry_titulo = tk.Entry(
    frame_form,
    font=FONT_MAIN,
    bg=BG_FIELD,
    fg=FG_WHITE,
    insertbackground="white",
    bd=1,
)
entry_titulo.pack(fill="x", ipady=4, pady=(2, 6))

tk.Label(
    frame_form,
    text="Autor / Setor:",
    font=FONT_LABEL,
    bg=BG_DARK,
    fg=FG_WHITE,
    anchor="w",
).pack(fill="x")
entry_autor = tk.Entry(
    frame_form,
    font=FONT_MAIN,
    bg=BG_FIELD,
    fg=FG_WHITE,
    insertbackground="white",
    bd=1,
)
entry_autor.pack(fill="x", ipady=4, pady=(2, 6))

tk.Label(
    frame_form,
    text="Resumo das Atividades:",
    font=FONT_LABEL,
    bg=BG_DARK,
    fg=FG_WHITE,
    anchor="w",
).pack(fill="x")
text_resumo = tk.Text(
    frame_form,
    height=3,
    font=FONT_MAIN,
    bg=BG_FIELD,
    fg=FG_WHITE,
    insertbackground="white",
    bd=1,
)
text_resumo.pack(fill="x", pady=(2, 6))

tk.Label(
    frame_form,
    text="Adicionar Tópico / Item:",
    font=FONT_LABEL,
    bg=BG_DARK,
    fg=FG_WHITE,
    anchor="w",
).pack(fill="x")
entry_item = tk.Entry(
    frame_form,
    font=FONT_MAIN,
    bg=BG_FIELD,
    fg=FG_WHITE,
    insertbackground="white",
    bd=1,
)
entry_item.pack(fill="x", ipady=4, pady=(2, 5))

frame_btn_item = tk.Frame(frame_form, bg=BG_DARK)
frame_btn_item.pack(fill="x", pady=2)

tk.Button(
    frame_btn_item,
    text="Adicionar",
    command=adicionar_item,
    bg="#3c3c3c",
    fg="white",
    bd=0,
    padx=10,
    pady=3,
    cursor="hand2",
).pack(side="left")
tk.Button(
    frame_btn_item,
    text="Remover",
    command=remover_item,
    bg="#b22222",
    fg="white",
    bd=0,
    padx=10,
    pady=3,
    cursor="hand2",
).pack(side="left", padx=5)

lista_itens = tk.Listbox(
    frame_form,
    height=3,
    font=FONT_MAIN,
    bg=BG_FIELD,
    fg=FG_WHITE,
    selectbackground="#007acc",
    bd=1,
)
lista_itens.pack(fill="x", pady=(5, 8))

btn_gerar_manual = tk.Button(
    janela,
    text="Gerar Relatório Manual (.docx)",
    command=gerar_manual,
    bg="#007acc",
    fg="white",
    font=FONT_LABEL,
    bd=0,
    padx=15,
    pady=8,
    cursor="hand2",
)
btn_gerar_manual.pack(pady=5)

lbl_status = tk.Label(
    janela, text="", font=("Segoe UI", 8, "italic"), bg=BG_DARK, fg="#aaaaaa"
)
lbl_status.pack(pady=(5, 2))

barra_progresso = ttk.Progressbar(
    janela,
    orient="horizontal",
    length=450,
    mode="determinate",
    style="Custom.Horizontal.TProgressbar",
)
barra_progresso.pack(pady=(0, 10))

janela.mainloop()